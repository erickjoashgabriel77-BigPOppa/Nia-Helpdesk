import os
import io
import streamlit as st
import streamlit.components.v1 as components
from utils import save_local_data

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def get_docx_template_path():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for ext in ["docx", "DOCX"]:
        path = os.path.join(base_dir, "Media", f"IT_Service_Ticket_Template.{ext}")
        if os.path.exists(path):
            return path
    return None


def generate_filled_docx(t):
    template_path = get_docx_template_path()
    if not template_path or not HAS_DOCX:
        return None

    client_name = t.get("client_name") or t.get("employee_name") or t.get("name") or "N/A"
    t_id = t.get("id", "N/A")
    date_created = t.get("date_created", "N/A")
    section = t.get("section", "N/A")
    unit = t.get("unit", "N/A")
    position = t.get("position", "N/A")
    supervisor = t.get("supervisor") or t.get("supervisor_name", "N/A")
    req_type = t.get("request_type") or t.get("category", "N/A")
    urgency = t.get("urgency", "N/A")
    status = t.get("status", "Pending")
    description = t.get("description", "N/A")
    equipment_type = t.get("equipment_type") or t.get("equipment", "N/A")
    serial_number = t.get("serial_number") or t.get("serial_no", "N/A")
    action_taken = t.get("action_taken", "N/A")
    serviced_by = t.get("serviced_by", "N/A")

    context = {
        "name": client_name,
        "client_name": client_name,
        "employee_name": client_name,
        "id": t_id,
        "ticket_id": t_id,
        "control_no": t_id,
        "date": date_created,
        "date_created": date_created,
        "section": section,
        "unit": unit,
        "position": position,
        "supervisor": supervisor,
        "supervisor_name": supervisor,
        "request_type": req_type,
        "category": req_type,
        "urgency": urgency,
        "status": status,
        "description": description,
        "issue": description,
        "equipment_type": equipment_type,
        "equipment": equipment_type,
        "serial_number": serial_number,
        "serial_no": serial_number,
        "sn": serial_number,
        "action_taken": action_taken,
        "serviced_by": serviced_by,
    }

    doc = Document(template_path)

    def replace_tokens(p):
        for key, val in context.items():
            tag = f"{{{{{key}}}}}"
            if tag in p.text:
                p.text = p.text.replace(tag, str(val))

    for p in doc.paragraphs:
        replace_tokens(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_tokens(p)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_ticket_slip_html(t):
    header_b64 = st.session_state.get("slip_header_b64", "")
    footer_b64 = st.session_state.get("slip_footer_b64", "")

    header_img = f'<img src="data:image/png;base64,{header_b64}" style="width:100%; max-height:100px; object-fit:contain; margin-bottom:10px;" />' if header_b64 else ''
    footer_img = f'<img src="data:image/png;base64,{footer_b64}" style="width:100%; max-height:70px; object-fit:contain; margin-top:15px;" />' if footer_b64 else ''

    client_name = t.get("client_name") or t.get("employee_name") or t.get("name") or "N/A"
    status = t.get("status", "Pending")
    equipment_type = t.get("equipment_type") or t.get("equipment", "N/A")
    serial_number = t.get("serial_number") or t.get("serial_no", "N/A")
    supervisor = t.get("supervisor") or t.get("supervisor_name", "N/A")

    if status == "Completed":
        badge_bg, badge_font = "#15803d", "#ffffff"
    elif status == "In Progress":
        badge_bg, badge_font = "#f59e0b", "#000000"
    elif status == "Cancelled":
        badge_bg, badge_font = "#b91c1c", "#ffffff"
    else:
        badge_bg, badge_font = "#e2e8f0", "#000000"

    return f"""
    <!DOCTYPE html>
    <html>
    <head>
        <style>
            body {{
                font-family: Arial, sans-serif;
                color: #0f172a;
                background: #ffffff;
                padding: 10px;
                margin: 0;
            }}
            .slip-container {{
                border: 2px solid #0f172a;
                padding: 15px;
                max-width: 600px;
                margin: 0 auto;
                background: #ffffff;
                border-radius: 6px;
                color: #000000;
            }}
            .slip-table {{
                width: 100%;
                border-collapse: collapse;
                margin-top: 10px;
            }}
            .slip-table th {{
                background-color: #f1f5f9;
                color: #000000;
                border: 1px solid #334155;
                padding: 6px 10px;
                text-align: left;
                font-size: 13px;
                font-weight: bold;
                width: 32%;
            }}
            .slip-table td {{
                border: 1px solid #334155;
                padding: 6px 10px;
                text-align: left;
                font-size: 13px;
                color: #000000;
                background-color: #ffffff;
            }}
            .header-bar {{
                background-color: #0f172a;
                color: #ffffff;
                padding: 8px;
                text-align: center;
                margin: 5px 0 12px 0;
                font-size: 15px;
                font-weight: bold;
                text-transform: uppercase;
                border-radius: 4px;
            }}
            .status-badge {{
                background-color: {badge_bg};
                color: {badge_font};
                padding: 3px 8px;
                border-radius: 4px;
                font-weight: bold;
                display: inline-block;
            }}
            .print-btn {{
                background-color: #0f172a;
                color: #ffffff;
                padding: 8px 16px;
                border: none;
                border-radius: 4px;
                cursor: pointer;
                font-weight: bold;
                margin-bottom: 12px;
            }}
            .print-btn:hover {{
                background-color: #1e293b;
            }}
            @media print {{
                .print-btn {{ display: none; }}
                body {{ padding: 0; background: #fff; }}
                .slip-container {{ border: 1px solid #000; }}
            }}
        </style>
    </head>
    <body>
        <button class="print-btn" onclick="window.print()">🖨️ Print / Save as PDF</button>
        <div class="slip-container">
            {header_img}
            <div class="header-bar">IT Support Service Request Slip</div>
            <table class="slip-table">
                <tr><th>Control No.</th><td><strong>{t.get('id', 'N/A')}</strong></td></tr>
                <tr><th>Date Submitted</th><td>{t.get('date_created', 'N/A')}</td></tr>
                <tr><th>Client Name</th><td>{client_name}</td></tr>
                <tr><th>Section & Unit</th><td>{t.get('section', 'N/A')} - {t.get('unit', 'N/A')}</td></tr>
                <tr><th>Position</th><td>{t.get('position', 'N/A')}</td></tr>
                <tr><th>Supervisor Name</th><td>{supervisor}</td></tr>
                <tr><th>Request Type</th><td>{t.get('request_type') or t.get('category', 'N/A')}</td></tr>
                <tr><th>Equipment Type</th><td>{equipment_type}</td></tr>
                <tr><th>Serial Number</th><td>{serial_number}</td></tr>
                <tr><th>Urgency Level</th><td>{t.get('urgency', 'N/A')}</td></tr>
                <tr><th>Status</th><td><span class="status-badge">{status}</span></td></tr>
                <tr><th>Description</th><td>{t.get('description', 'N/A')}</td></tr>
            </table>
            {footer_img}
        </div>
    </body>
    </html>
    """


def render_admin_views(admin_menu):
    tickets = st.session_state.get("tickets", [])

    if admin_menu == "📊 Ticket Management Dashboard":
        st.header("📊 IT Support Ticket Dashboard")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Tickets", len(tickets))
        c2.metric("Pending", len([t for t in tickets if t.get("status") == "Pending"]))
        c3.metric("In Progress", len([t for t in tickets if t.get("status") == "In Progress"]))
        c4.metric("Completed", len([t for t in tickets if t.get("status") == "Completed"]))

        st.markdown("---")

        if not tickets:
            st.info("No tickets recorded in the system.")
            return

        status_options = ["Pending", "In Progress", "Completed", "Cancelled"]

        for t in reversed(tickets):
            t_id = t.get("id", "N/A")
            client_name = t.get("client_name") or t.get("employee_name") or t.get("name") or "Unknown Client"
            section = t.get("section", "N/A")
            status = t.get("status", "Pending")
            request_type = t.get("request_type") or t.get("category", "N/A")
            unit = t.get("unit", "N/A")
            position = t.get("position", "N/A")
            supervisor = t.get("supervisor") or t.get("supervisor_name", "N/A")
            date_created = t.get("date_created", "N/A")
            urgency = t.get("urgency", "N/A")
            description = t.get("description", "N/A")
            equipment_type = t.get("equipment_type") or t.get("equipment", "N/A")
            serial_number = t.get("serial_number") or t.get("serial_no", "N/A")

            status_index = status_options.index(status) if status in status_options else 0

            with st.expander(f"🎫 **{t_id}** | {client_name} ({section}) - Status: {status}"):
                col1, col2 = st.columns(2)
                with col1:
                    st.write(f"**Category:** {request_type}")
                    st.write(f"**Unit:** {unit}")
                    st.write(f"**Position:** {position}")
                    st.write(f"**Supervisor:** {supervisor}")
                    st.write(f"**Equipment Type:** {equipment_type}")
                    st.write(f"**Serial No.:** {serial_number}")
                    st.write(f"**Date Submitted:** {date_created}")
                with col2:
                    st.write(f"**Urgency:** {urgency}")
                    st.write(f"**Description:** {description}")

                st.markdown("---")
                act_col1, act_col2 = st.columns([2, 2])

                with act_col1:
                    new_status = st.selectbox(
                        "Update Status",
                        status_options,
                        index=status_index,
                        key=f"status_select_{t_id}",
                    )
                    if st.button("Update Ticket Status", key=f"btn_update_{t_id}"):
                        t["status"] = new_status
                        save_local_data()
                        st.success(f"Status for {t_id} updated to {new_status}!")
                        st.rerun()

                with act_col2:
                    st.write("**Printable & Downloadable Slip**")
                    b_col1, b_col2 = st.columns(2)

                    with b_col1:
                        with st.popover("🖨️ View Slip", use_container_width=True):
                            slip_html = generate_ticket_slip_html(t)
                            components.html(slip_html, height=520, scrolling=True)

                    with b_col2:
                        docx_bytes = generate_filled_docx(t)
                        if docx_bytes:
                            st.download_button(
                                label="📄 Word Slip",
                                data=docx_bytes,
                                file_name=f"Ticket_{t_id}.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True,
                                key=f"dl_docx_{t_id}",
                            )
                        elif not HAS_DOCX:
                            st.caption("⚠️ Add `python-docx` to requirements.txt")
                        else:
                            st.caption("⚠️ Template file missing")

    elif admin_menu == "👥 Employee / Client Management":
        st.header("👥 Client / Employee Directory")
        employees = st.session_state.get("employees", [])

        st.dataframe(employees, use_container_width=True)

        with st.expander("➕ Register New Employee"):
            with st.form("add_emp_form"):
                emp_id = st.text_input("Employee ID", placeholder="EMP-EOS-010")
                emp_name = st.text_input("Full Name")
                section = st.selectbox("Section", ["Engineering & Operations Section (EOS)", "Administrative & Finance Section (AFS)"])
                unit = st.text_input("Unit")

                if st.form_submit_button("Add Employee"):
                    if emp_id and emp_name:
                        st.session_state.employees.append({"id": emp_id, "name": emp_name, "section": section, "unit": unit})
                        save_local_data()
                        st.success(f"Employee {emp_name} registered successfully!")
                        st.rerun()

    elif admin_menu == "🔑 IT Admin Account Management":
        st.header("🔑 IT Admin Passkey & Account Management")

        st.write("Current Admin Passkey configured in system.")
        new_key = st.text_input("Change Admin Passkey:", value=st.session_state.get("admin_passkey", "ADMIN2026"), type="password")

        if st.button("Save Passkey"):
            st.session_state.admin_passkey = new_key
            save_local_data()
            st.success("Admin passkey updated successfully!")